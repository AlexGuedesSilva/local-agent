"""Fill the Lindt integration specification template with IAC-backed facts."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject

from agent.documentation.skyone_iac import analyze_skyone_iac, module_parameters, plain_text


def _set_paragraph(paragraph: Any, value: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def _set_cell(table: Any, row: int, column: int, value: str) -> None:
    cell = table.cell(row, column)
    _set_paragraph(cell.paragraphs[0], value)
    for paragraph in cell.paragraphs[1:]:
        _set_paragraph(paragraph, "")


def _safe_url(value: Any) -> str:
    rendered = str(value)
    if "@" in rendered and "://" in rendered:
        scheme, remainder = rendered.split("://", 1)
        if "@" in remainder:
            rendered = f"{scheme}://[credenciais removidas]@{remainder.split('@', 1)[1]}"
    return rendered


def _return_status(modules: list[dict[str, Any]], status: str) -> str:
    for module in modules:
        if module.get("subtype") != "return":
            continue
        parameters = module_parameters(module)
        if str(parameters.get("integra_configuration_status_code", "")) == status:
            body = parameters.get("jsonata", "")
            return str(body) if body else str(module.get("name", "Retorno configurado"))
    return "Não configurado no IAC."


def _fill_document(document: DocumentObject, modules: list[dict[str, Any]]) -> None:
    if len(document.tables) < 14:
        raise ValueError(
            "O modelo não corresponde ao Modelo_Especificacao_Integracao_Lindt esperado."
        )
    facts = analyze_skyone_iac(modules)
    tables = document.tables
    endpoint = _safe_url(facts["endpoint"])
    source = str(facts["source_hint"])
    destination = str(facts["destination_hint"])
    today = date.today().strftime("%d/%m/%Y")

    identity = [
        "Não informado no IAC",
        "Cadastro de produto/material com geração de XML e envio SFTP (validar nome oficial)",
        "Não informado no IAC",
        f"{source} (inferido pelo nome do payload; confirmar sistema de origem)",
        f"{destination} (identificação indicada pelo conector e pelos retornos; confirmar)",
        "Não informado no IAC",
        "Não informado no IAC",
        "Não informado no IAC",
        today,
        "1.0",
    ]
    for row, value in enumerate(identity):
        _set_cell(tables[1], row, 1, value)

    _set_cell(tables[2], 0, 1, "( ) Temporal/Agendado    (X) Evento/Webhook    ( ) Outro")
    for row in range(len(tables[3].rows)):
        _set_cell(tables[3], row, 1, "Não se aplica: o gatilho exportado é API Gateway.")

    trigger = facts["trigger"]
    trigger_name = str(trigger.get("name", "API Gateway"))
    event_values = [
        f"Sistema chamador não identificado no IAC (gatilho: {trigger_name}).",
        f"{endpoint} (rota de QA indicada no IAC).",
        "Não informado no IAC.",
        "Não informado no IAC.",
        "Requisição HTTP com dados de produto/material; descrição do gatilho menciona entrega de saída, confirmar finalidade.",
        "Estrutura acessada: body.MT_PRODUCT_MASTER_DATA_SAP_REQ.products.product ou body.n0:MT_PRODUCT_MASTER_DATA_SAP_REQ.products.product. O IAC não contém valores de exemplo.",
    ]
    for row, value in enumerate(event_values):
        _set_cell(tables[4], row, 1, value)

    source_values = [
        f"{source} (inferência baseada em MT_PRODUCT_MASTER_DATA_SAP_REQ; confirmar).",
        "Não informado; o endpoint contém /qa.",
        "HTTP/API Gateway.",
        f"{endpoint}.",
        "Não informado no IAC.",
        "Não informado no IAC.",
        "Não informado no IAC.",
        "Endpoint de QA identificado; ambiente da origem não confirmado.",
        "Não informado no IAC.",
    ]
    for row, value in enumerate(source_values):
        _set_cell(tables[5], row, 1, value)

    input_rows = [
        (
            "MT_PRODUCT_MASTER_DATA_SAP_REQ (ou n0:MT_PRODUCT_MASTER_DATA_SAP_REQ)",
            "Objeto (estrutura inferida pelo acesso no JSONata)",
            "Sim (usado pela transformação)",
            "Nó raiz do payload aceito pelo fluxo, com ou sem prefixo n0.",
            "Não incluído no IAC.",
        ),
        (
            "products.product",
            "Não informado no IAC",
            "Sim (condição explícita do fluxo)",
            "Identificador de produto/material; deve existir e não estar vazio.",
            "Não incluído no IAC.",
        ),
    ]
    for row, values in enumerate(input_rows, start=1):
        for column, value in enumerate(values):
            _set_cell(tables[6], row, column, value)
    for row in range(1 + len(input_rows), len(tables[6].rows)):
        for column in range(len(tables[6].columns)):
            _set_cell(tables[6], row, column, "")

    _set_paragraph(
        document.paragraphs[27],
        "O corpo deve conter MT_PRODUCT_MASTER_DATA_SAP_REQ (com ou sem o prefixo n0) e products.product deve existir e ser diferente de vazio.",
    )
    _set_paragraph(
        document.paragraphs[28],
        "A validação não comprova outros campos obrigatórios nem valida o schema completo do produto.",
    )
    _set_paragraph(document.paragraphs[30], "Nenhum filtro de exclusão está definido no IAC.")
    _set_paragraph(
        document.paragraphs[32],
        "Gera fileName como MD_<product>_<AAAA><MM><DD>_<hh><mm><ss>.xml, com deslocamento de horário -0300. Converte o objeto JSON em XML e acrescenta os namespaces configurados.",
    )
    _set_paragraph(document.paragraphs[34], "Tratamento de duplicidade não informado no IAC.")
    _set_paragraph(
        document.paragraphs[36],
        "Se a validação falhar, registra log e retorna HTTP 400. Se createdFile não existir após a gravação, registra log e retorna HTTP 500. Retentativas, timeout e idempotência não estão especificados.",
    )

    target_values = [
        f"{destination} (o IAC cita Nexxa FTP e também DHL Chile; confirmar destino oficial).",
        "Não informado no IAC.",
        "SFTP/conector de arquivo.",
        "Host e diretório remoto não incluídos no IAC.",
        "Não informado no IAC.",
        "Não informado no IAC.",
        "Não informado no IAC.",
        "Não informado no IAC.",
        "O IAC só permite verificar a existência da saída createdFile; detalhes operacionais não aparecem no export.",
    ]
    for row, value in enumerate(target_values):
        _set_cell(tables[8], row, 1, value)

    output_values = [
        "XML.",
        "JSON de entrada transformado para o nó n0:MT_PRODUCT_MASTER_DATA_SAP_REQ com namespaces XML. O IAC não inclui um XML de exemplo.",
        "products.product é a única propriedade de negócio explicitamente validada no fluxo.",
        "Arquivo MD_<product>_<AAAA><MM><DD>_<hh><mm><ss>.xml; timestamp com deslocamento -0300; detalhes de encoding não informados.",
        "Não informado no IAC.",
    ]
    for row, value in enumerate(output_values):
        _set_cell(tables[9], row, 1, value)

    mappings = [
        (
            "MT_PRODUCT_MASTER_DATA_SAP_REQ",
            "n0:MT_PRODUCT_MASTER_DATA_SAP_REQ",
            "Direto + constante",
            "Preserva o objeto raiz e acrescenta xmlns:n0 e xmlns:prx.",
            "Exemplo de payload não fornecido.",
        ),
        (
            "products.product",
            "fileName",
            "Concatenação",
            "Concatena prefixo MD_, código do produto, data/hora e extensão .xml.",
            "Código e valor de exemplo não fornecidos.",
        ),
        (
            "documento (objeto JSON preparado)",
            "xml (arquivo enviado)",
            "Conversão",
            "O módulo JSON_TO_XML converte o objeto usando $ para atributo e _ para caractere customizado.",
            "XML de exemplo não fornecido.",
        ),
        (
            "fileName gerado",
            "Nome do arquivo no SFTP",
            "Direto",
            "O conector recebe fileName e o conteúdo xml.",
            "Nome final depende do produto e do timestamp.",
        ),
    ]
    for row, values in enumerate(mappings, start=1):
        for column, value in enumerate(values):
            _set_cell(tables[11], row, column, value)

    for column, value in enumerate(["1.0", today, "Não informado", "Preenchimento inicial com base no export IAC"]):
        _set_cell(tables[12], 1, column, value)
    for row in range(1, len(tables[13].rows)):
        for column in (1, 2, 3):
            _set_cell(tables[13], row, column, "Pendente de aprovação")

    _set_paragraph(
        document.paragraphs[37],
        "Pontos para validação: "
        + " ".join(plain_text(warning) for warning in facts["warnings"])
        + " Método HTTP, autenticação, schema completo, host SFTP, owners, criticidade, retentativas e política de duplicidade também precisam ser confirmados.",
    )


def create_filled_lindt_template(
    template_path: str,
    modules: list[dict[str, Any]],
) -> bytes:
    """Return a filled copy of the supplied Lindt DOCX template as bytes."""
    document = Document(template_path)
    _fill_document(document, modules)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
